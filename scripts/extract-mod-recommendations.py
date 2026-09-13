"""Read-only extraction from the MTC workbook. Never exports/edits the workbook.

Usage: bundled-python scripts/extract-mod-recommendations.py work/mod-tree-source.xlsx
Only catalogue/formula evidence is retained, never the workbook's player inputs.
"""
import hashlib
import json
from pathlib import Path
import sys
import openpyxl

root = Path(__file__).resolve().parents[1]
path = Path(sys.argv[1])
formulas = openpyxl.load_workbook(path, data_only=False)
values = openpyxl.load_workbook(path, data_only=True)
reference = json.loads((root / "lib/cifi/mod-tree/reference.json").read_text(encoding="utf-8"))
codes = {row["code"]: row for row in reference["nodes"]}


def scalar(value):
    # Sheets' log -> number roundtrip leaves e.g. 39.99999999999999 for 40.
    # Imported constants carry at most 15 meaningful spreadsheet digits.
    if isinstance(value, (int, float)):
        return format(value, ".15g")
    if isinstance(value, str) and not value.startswith("#"):
        return value
    raise ValueError(f"Missing/error source constant: {value!r}")


def expression(cell):
    value = cell.value
    if hasattr(value, "text"):
        assert value.ref == cell.coordinate, "Unexpected spilling effect formula"
        return value.text
    if cell.data_type == "f":
        return value
    return float(scalar(value))


items = []
ref = values["ModRef Local"]
calc = formulas["ModCalc Local"]
for row in range(2, 276):
    code = ref.cell(row, 2).value
    known = codes[code]
    assert ref.cell(row, 1).value == known["name"]
    assert calc.cell(row, 1).value == known["name"]
    assert values["ModRef Import"].cell(row, 2).value == code
    prerequisites = [s.strip() for s in (values["ModRef Import"].cell(row, 3).value or "").split(",") if s.strip()]
    assert prerequisites == known["prerequisites"]
    factors = []
    for col in (9, 10):
        value = expression(formulas["ModRef Local"].cell(row, col))
        if isinstance(value, str):
            value = value.replace("MOD_CODE_TO_LEVEL_NF(INDIRECT(ModRef_LocalModCodes))", f'MOD_CODE_TO_LEVEL_NF("{code}")')
        factors.append(value)
    items.append({"code": code, "sourceRow": row, "startCost": scalar(ref.cell(row, 7).value), "costGrowth": scalar(ref.cell(row, 8).value),
                  "startBase": factors[0], "baseGrowth": factors[1], "effects": [expression(calc.cell(row, c)) for c in range(2, 42)]})
assert len(items) == len(codes) == 274
source = {
    "url": reference["source"], "version": values["FAQ & Credits"]["O3"].value,
    "retrieved": "2026-09-12", "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    "costRange": "ModRef Local!G2:J275", "effectRange": "ModCalc Local!B2:AO275", "powerCell": "ModCalc Local!AP2", "unlockCell": "ModRef Local!C2",
    "numericPolicy": "Source constants normalized to 15 significant digits; integer costs and MP subtraction use BigInt. Effect powers use overflow-safe logarithms.",
    "nodes": items,
}
(root / "lib/cifi/mod-tree/recommendation-data.json").write_text(json.dumps(source, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
audit = {"source": {k: v for k, v in source.items() if k != "nodes"}, "powerFormula": formulas["ModCalc Local"]["AP2"].value.text,
         "costFormula": formulas["ModRef Local"]["F2"].value.text, "scoreFormula": formulas["Mods"]["E30"].value.text,
         "powerRange": values["Dynamic Ranges"]["D15"].value, "subLogDefinition": formulas.defined_names["SUB_LOG_NF"].attr_text}
(root / "docs/mod-tree-recommendation-source.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
# This public source snapshot is a blank starting profile. Refuse to export any
# fixture if real player levels/progress were populated in a future source.
assert all(not ref.cell(r, 4).value for r in range(2, 276))
assert all(not values["ModValues"].cell(r, 8).value for r in range(4, 24))
assert all(not values["ModValues"].cell(r, 12).value for r in range(4, 18))
fixture = {"sourceSha256": source["sha256"], "profile": "all node levels and progress zero; weights 1/12/10/8/24/72/6/1",
           "powerLog": {ref.cell(r, 2).value: values["ModCalc Local"].cell(r, 43).value for r in range(2, 276)}}
target = root / "tests/fixtures/mod-tree-sheet-zero.json"
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(fixture, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(f"Extracted {len(items)} nodes and {sum(isinstance(v, str) for r in items for v in r['effects'])} effect formulas. SHA256 {source['sha256']}")
